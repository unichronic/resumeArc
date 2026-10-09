#!/usr/bin/env python3
"""Build a Google-Sheets-friendly internship tracker from local notes."""

from __future__ import annotations

import csv
import re
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parents[1]
OUT_XLSX = ROOT / "internship_tracker_2026_05_27.xlsx"
OUT_CSV = ROOT / "internship_tracker_2026_05_27_opportunities.csv"

STATUS_FILE = ROOT / "unapplied_jobs_draft_2026_05_26.md"

OPPORTUNITY_TABLE_FILES = [
    ROOT / "internship_opportunities_2026-05-13.md",
    ROOT / "internship_opportunities_last_30_days_2026-05-13.md",
    ROOT / "tailored_resumes/recent_interns/resume_fit_report.md",
    ROOT / "tailored_resumes/selected_internships/resume_fit_report.md",
    ROOT / "tailored_resumes/next_applications_2026_05_25/resume_fit_report.md",
    ROOT / "tailored_resumes/email_leads_2026_05_25/resume_fit_report.md",
    ROOT / "tailored_resumes/social_leads_2026_05_26/resume_fit_report.md",
    ROOT / "next_application_materials.md",
    ROOT / "email_lead_application_materials.md",
    ROOT / "social_lead_application_materials.md",
]

DRAFT_FILES = [
    ROOT / "selected_internship_email_drafts.md",
    ROOT / "next_application_materials.md",
    ROOT / "email_lead_application_materials.md",
    ROOT / "social_lead_application_materials.md",
    ROOT / "active_email_cold_mails_2026_05_26.md",
    ROOT / "metronis_backend_infra_outreach.md",
    ROOT / "subimage_followup.md",
]
DRAFT_FILES.extend(sorted(ROOT.glob("cold_email_drafts_batch_*.md")))

CONTACT_FILES = [
    ROOT / "cold_mail_contacts_public.md",
    ROOT / "founder_contacts_public_2026-05-11.md",
    ROOT / "contact_to_field_research_2026-05-13.md",
    *DRAFT_FILES,
]

TARGET_RESEARCH_FILES = [
    ROOT / "outreach_category_recheck_2026-05-11.md",
    ROOT / "startup_outreach_tiers.md",
    ROOT / "startup_backend_infra_ai_tiers.md",
    ROOT / "backend_infra_ai_rankings.md",
]

EMAIL_RE = re.compile(r"[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}")
URL_RE = re.compile(r"https?://[^\s)`>]+")

COMPANY_ALIASES = {
    "banking stack": "CredResolve / Banking Stack",
    "credresolve": "CredResolve / Banking Stack",
    "credresolve abhijit das": "CredResolve / Banking Stack",
    "banking stack tech intern": "CredResolve / Banking Stack",
    "omli kids": "Omli",
    "omli ai/ml or platform": "Omli",
    "binary": "Binary form lead",
    "proofbuild": "proof_of_Build",
    "proof_of_build": "proof_of_Build",
    "proof of build": "proof_of_Build",
    "cloudsek": "CloudSEK",
    "flywire": "FlyWire",
    "tidb": "TiDB / PingCAP",
    "pingcap": "TiDB / PingCAP",
    "ibm cloudability": "IBM Cloudability / Apptio",
    "cloudthat technologies": "CloudThat",
    "cloudthat technologies pvt ltd": "CloudThat",
    "i2k2 networks pvt limited": "i2k2 Networks",
    "i2k2": "i2k2 Networks",
    "miniorange": "miniOrange",
    "allen": "ALLEN Digital",
    "allendigital": "ALLEN Digital",
    "testmu": "TestMu AI",
    "testmu ai": "TestMu AI",
    "neosapien": "NeoSapien",
    "navi": "Navi",
    "walmart": "Walmart Global Tech",
    "walmart global tech": "Walmart Global Tech",
    "dafi labs": "Dafi Labs",
    "factech": "Factech",
    "ripik": "Ripik.ai",
    "ripik.ai": "Ripik.ai",
    "statsby": "Statsby Solutions",
    "statsby solutions": "Statsby Solutions",
    "morphle": "Morphle Labs",
    "morphle labs": "Morphle Labs",
    "sigtuple": "SigTuple Technologies",
    "sigtuple technologies": "SigTuple Technologies",
    "deckoviz": "Deckoviz Space Labs",
    "deckoviz space labs": "Deckoviz Space Labs",
    "fusionpact": "Fusionpact Technologies",
    "fusionpact technologies": "Fusionpact Technologies",
    "bigfig ai": "BigFig.AI",
    "xyrs clone": "XYRS Clone",
}

KNOWN_SINGLE_COMPANY_ROLES = {
    "AICTE Python Full Stack Internship": (
        "AICTE Internship Portal",
        "Python Full Stack Internship",
    ),
    "CertifiedAIJobs - AI Software Tester": (
        "CertifiedAIJobs",
        "AI Software Tester",
    ),
    "NIELIT Summer Internship": (
        "NIELIT",
        "Summer Internship",
    ),
}

HEADER_FILL = PatternFill("solid", fgColor="1F2937")
SUBHEADER_FILL = PatternFill("solid", fgColor="E5E7EB")
ODD_FILL = PatternFill("solid", fgColor="FFFFFF")
EVEN_FILL = PatternFill("solid", fgColor="F8FAFC")
THIN_BORDER = Border(bottom=Side(style="thin", color="CBD5E1"))

STATUS_FILLS = {
    "Open": "DCFCE7",
    "Applied by draft": "E5E7EB",
    "Needs verification": "FEF3C7",
    "Eliminated": "FEE2E2",
    "Closed / inactive": "F3F4F6",
    "Skip": "E5E7EB",
    "Unclassified": "F1F5F9",
}

PRIORITY_FILLS = {
    "P0": "FEE2E2",
    "P0 if eligible": "FFE4E6",
    "P1": "FEF3C7",
    "P1 if eligible": "FFF7ED",
    "P2": "DBEAFE",
    "P2 if eligible": "E0F2FE",
    "1": "DCFCE7",
    "2": "DCFCE7",
    "3": "FEF3C7",
    "4": "FEF3C7",
    "5": "DBEAFE",
    "6": "DBEAFE",
    "7": "DBEAFE",
    "8": "F3F4F6",
    "9": "F3F4F6",
}


def rel(path: Path | str) -> str:
    path = Path(path)
    try:
        return path.resolve().relative_to(ROOT).as_posix()
    except Exception:
        return str(path)


def clean_md(text: str | None) -> str:
    if not text:
        return ""
    text = str(text)
    text = text.replace("<br>", " ").replace("<br/>", " ")
    text = re.sub(r"`([^`]*)`", r"\1", text)
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", text)
    text = text.replace("**", "").replace("__", "").replace("~~", "")
    text = text.replace("\\|", "|")
    return re.sub(r"\s+", " ", text).strip()


def norm_key(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", clean_md(text).lower())


def canonical_company(company: str) -> str:
    company = clean_md(company)
    company = re.sub(r"^\d+\.\s*", "", company)
    company = re.sub(r"\s+(Email|Note|DM|Form|Comment|WhatsApp)$", "", company, flags=re.I)
    base = company.strip(" :-")
    alias_key = re.sub(r"[^a-z0-9]+", " ", base.lower()).strip()
    compact = norm_key(base)
    return COMPANY_ALIASES.get(alias_key) or COMPANY_ALIASES.get(compact) or base


def canonical_role(company: str, role: str) -> str:
    company = canonical_company(company)
    role = clean_md(role) or "General / careers target"
    role = re.sub(r"\s*,?\s*Summer 2026$", "", role, flags=re.I)
    role = role.replace(" - ", ", ")
    role = role.replace("SDE - Backend - Intern", "SDE Backend Intern")
    if company == "Anakin" and "Software Engineering Intern" in role:
        return "Software Engineering Intern, Backend / Systems / Scraping Infra"
    if company == "NeoSapien" and "Backend Intern" in role:
        return "Backend Intern / AI Intern"
    if company == "Cekura" and "AI Engineer Intern" in role:
        return "AI Engineer Intern"
    if company == "Readily" and "Full-Stack Engineering Intern" in role:
        return "Full-Stack Engineering Intern"
    if company == "Promptless" and "Software Engineering Intern" in role:
        return "Software Engineering Intern"
    if company == "Reacher" and "Software Engineer Intern" in role:
        return "Software Engineer Intern"
    if company == "Peculiar Technologies" and "Software Development Intern" in role:
        return "Software Development Internship"
    if company == "proof_of_Build" and "AI + Crypto" in role:
        return "AI + Crypto Software Engineer Intern / Trading Agent"
    if company == "CredResolve / Banking Stack" and role == "General / careers target":
        return "Tech Intern"
    if company == "Dafi Labs" and "MERN / Blockchain" in role:
        return "MERN Stack Intern"
    if company == "Navi" and "CloudOps Engineer" in role:
        return "CloudOps Engineer I"
    if company == "Walmart Global Tech" and role.startswith("Grad Intern"):
        return "Grad Intern, No Work Experience"
    if company == "Pulse" and ("ML Engineer" in role or "Machine Learning Engineer" in role):
        return "Machine Learning Engineer Intern"
    if company == "SID" and "Research Intern" in role:
        return "Research Intern"
    return role


def split_company_role(label: str) -> tuple[str, str]:
    label = clean_md(label).strip()
    label = re.sub(r"^\d+\.\s*", "", label)
    if label in KNOWN_SINGLE_COMPANY_ROLES:
        return KNOWN_SINGLE_COMPANY_ROLES[label]
    for known, result in KNOWN_SINGLE_COMPANY_ROLES.items():
        if label.lower() == known.lower():
            return result
    for sep in [" - ", " – ", " — "]:
        if sep in label:
            company, role = label.split(sep, 1)
            company = canonical_company(company)
            return company, canonical_role(company, role)
    no_sep_patterns = [
        (r"^(.+?)\s+(Tech Intern)$", None),
        (r"^(.+?)\s+(Software Engineer Intern)$", None),
        (r"^(.+?)\s+(Python Backend and AI Intern)$", None),
        (r"^(.+?)\s+(AI Engineers Internship)$", None),
        (r"^(.+?)\s+(AI Intern)$", None),
        (r"^(.+?)\s+(CloudOps Engineer\s*-?\s*I)$", None),
        (r"^(.+?)\s+(Grad Intern)$", None),
        (r"^(.+?)\s+(Backend Intern)$", None),
        (r"^(.+?)\s+(Full Stack Developer Intern)$", None),
        (r"^(.+?)\s+(Full Stack Intern)$", None),
        (r"^(.+?)\s+(Machine Learning Intern)$", None),
        (r"^(.+?)\s+(ML Engineer Intern)$", None),
        (r"^(.+?)\s+(Research Intern)$", None),
    ]
    for pattern, _ in no_sep_patterns:
        match = re.match(pattern, label, flags=re.I)
        if match:
            company = canonical_company(match.group(1))
            return company, canonical_role(company, match.group(2))
    if label.lower() == "metronis foundation team":
        return "Metronis", "Foundation Team Intern"
    if label.lower().startswith("omli"):
        return "Omli", "Backend / Platform Intern"
    if label.lower().startswith("aicte "):
        return "AICTE Internship Portal", "Python Full Stack Internship"
    company = canonical_company(label)
    return company, canonical_role(company, "General / careers target")


def opportunity_key(company: str, role: str) -> str:
    company = canonical_company(company)
    role = canonical_role(company, role)
    return f"{norm_key(company)}|{norm_key(role)}"


def unique_join(values: list[str] | set[str]) -> str:
    seen = []
    for value in values:
        value = clean_md(value)
        if value and value not in seen:
            seen.append(value)
    return " | ".join(seen)


def append_unique(existing: str, value: str, sep: str = " | ") -> str:
    value = clean_md(value)
    if not value:
        return existing
    parts = [p.strip() for p in existing.split(sep) if p.strip()] if existing else []
    if value not in parts:
        parts.append(value)
    return sep.join(parts)


def infer_source_date(path: Path) -> str:
    match = re.search(r"20\d{2}[-_]\d{2}[-_]\d{2}", path.as_posix())
    if match:
        return match.group(0).replace("_", "-")
    return ""


def extract_links_emails(text: str) -> str:
    emails = [email for email in EMAIL_RE.findall(text) if email != "ishuvam.pal@gmail.com"]
    values = emails + URL_RE.findall(text)
    return unique_join(values)


def resolve_resume_path(cell: str, source_path: Path) -> str:
    cell = clean_md(cell)
    if not cell or cell.lower() == "none":
        return ""
    first = cell.split("|")[0].strip()
    if not first.lower().endswith(".pdf"):
        match = re.search(r"[\w./-]+\.pdf", cell)
        first = match.group(0) if match else first
    first = first.strip()
    if not first.lower().endswith(".pdf"):
        return cell
    candidate = ROOT / first
    if candidate.exists():
        return rel(candidate)
    candidate = source_path.parent / first
    if candidate.exists():
        return rel(candidate)
    return first


def parse_table_line(line: str) -> list[str]:
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [clean_md(part) for part in line.split("|")]


def iter_markdown_tables(path: Path):
    if not path.exists():
        return
    lines = path.read_text(encoding="utf-8", errors="replace").splitlines()
    heading = ""
    i = 0
    while i < len(lines):
        line = lines[i]
        heading_match = re.match(r"^(#{1,4})\s+(.+)$", line)
        if heading_match:
            heading = clean_md(heading_match.group(2))
        if (
            line.lstrip().startswith("|")
            and i + 1 < len(lines)
            and re.match(r"^\s*\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)+\|?\s*$", lines[i + 1])
        ):
            headers = parse_table_line(line)
            i += 2
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                cells = parse_table_line(lines[i])
                row = dict(zip(headers, cells))
                yield heading, row
                i += 1
            continue
        i += 1


def default_row(company: str, role: str) -> dict:
    company = canonical_company(company)
    role = canonical_role(company, role)
    return {
        "Status": "Unclassified",
        "Priority": "",
        "Company": canonical_company(company),
        "Role": clean_md(role),
        "Category": "",
        "Fit": "",
        "Location / Constraint": "",
        "Apply Route": "",
        "Apply Link / Email": "",
        "Resume PDF": "",
        "Attachment PDF": "",
        "Materials / Draft File": "",
        "Source Files": "",
        "Notes": "",
        "Last Seen": "",
        "Next Action": "",
    }


def upsert_opportunity(opps: dict[str, dict], company: str, role: str, source: Path | None = None, **fields) -> dict:
    company = canonical_company(company)
    role = canonical_role(company, role)
    key = opportunity_key(company, role)
    row = opps.setdefault(key, default_row(company, role))
    if source:
        row["Source Files"] = append_unique(row["Source Files"], rel(source))
        row["Last Seen"] = row["Last Seen"] or infer_source_date(source)
    for field, value in fields.items():
        field = {
            "Last_Seen": "Last Seen",
            "Location__Constraint": "Location / Constraint",
        }.get(field, field)
        if field not in row:
            continue
        value = clean_md(value)
        if not value:
            continue
        if field == "Status":
            precedence = {
                "": 0,
                "Unclassified": 0,
                "Open": 1,
                "Needs verification": 2,
                "Closed / inactive": 3,
                "Eliminated": 3,
                "Skip": 3,
                "Applied by draft": 4,
            }
            if precedence.get(value, 0) > precedence.get(row["Status"], 0):
                row["Status"] = value
        elif field in {"Notes", "Source Files", "Materials / Draft File", "Apply Link / Email"}:
            row[field] = append_unique(row[field], value)
        elif field == "Resume PDF":
            existing = row[field]
            if not existing or "omli_ai_ml_platform" in existing:
                row[field] = value
            elif "/" in value and "/" not in existing.split(" | ")[0]:
                row[field] = value
            elif value not in row[field]:
                row[field] = append_unique(row[field], value)
        elif not row[field]:
            row[field] = value
        elif value != row[field] and field in {"Fit", "Apply Route", "Location / Constraint", "Category"}:
            row[field] = append_unique(row[field], value)
    return row


def parse_status_file() -> dict[str, dict]:
    opps: dict[str, dict] = {}
    if not STATUS_FILE.exists():
        return opps
    section = ""
    for line in STATUS_FILE.read_text(encoding="utf-8", errors="replace").splitlines():
        section_match = re.match(r"^##\s+(.+)$", line)
        if section_match:
            section = clean_md(section_match.group(1))
            continue
        match = re.match(r"^-\s+\[( |x)\]\s+(.+)$", line)
        if not match:
            continue
        checked = match.group(1) == "x"
        text = clean_md(match.group(2))
        main = text
        note = ""
        note_match = re.search(r"\.\s+(User|File|Mentioned|Do not|LinkedIn|Listing|I could|The two).+", text)
        if note_match:
            main = text[: note_match.start()].strip()
            note = text[note_match.start() + 2 :].strip()
        company, role = split_company_role(main)
        lower_text = text.lower()
        if checked:
            status = "Applied by draft"
        elif "low confidence" in section.lower() or "eliminated" in section.lower():
            if "skip" in lower_text or "paid certification" in lower_text:
                status = "Skip"
            elif "no longer accepting" in lower_text or "unavailable" in lower_text:
                status = "Closed / inactive"
            elif "eliminated" in lower_text or "poor fit" in lower_text:
                status = "Eliminated"
            else:
                status = "Needs verification"
        else:
            status = "Open"
        category = (
            "Cold outreach target"
            if "cold-outreach" in section.lower()
            else "Needs verification / eliminated"
            if "low confidence" in section.lower()
            else "Specific internship lead"
        )
        upsert_opportunity(
            opps,
            company,
            role,
            STATUS_FILE,
            Status=status,
            Category=category,
            Notes=note,
            Last_Seen="2026-05-26",
        )
    return opps


def parse_opportunity_tables(opps: dict[str, dict]) -> None:
    for path in OPPORTUNITY_TABLE_FILES:
        if not path.exists():
            continue
        for heading, row in iter_markdown_tables(path):
            headers = {h.lower(): h for h in row}
            company = row.get(headers.get("company", ""), "")
            role = row.get(headers.get("role", ""), "")
            lead = row.get(headers.get("lead", ""), "")
            company_role = row.get(headers.get("company / role", ""), "")
            if not company and not role and not lead and not company_role:
                continue
            if company and role:
                parsed_company, parsed_role = canonical_company(company), role
            else:
                parsed_company, parsed_role = split_company_role(role or lead or company_role)
            priority = row.get(headers.get("priority", ""), "")
            fit = (
                row.get(headers.get("fit", ""), "")
                or row.get(headers.get("fit note", ""), "")
                or row.get(headers.get("fit / caveat", ""), "")
                or row.get(headers.get("why it fits you", ""), "")
                or row.get(headers.get("gap / caution", ""), "")
                or row.get(headers.get("note", ""), "")
            )
            location = (
                row.get(headers.get("location", ""), "")
                or row.get(headers.get("location / constraint", ""), "")
            )
            source = (
                row.get(headers.get("link", ""), "")
                or row.get(headers.get("source", ""), "")
                or row.get(headers.get("source type", ""), "")
            )
            apply_route = (
                row.get(headers.get("apply", ""), "")
                or row.get(headers.get("apply link", ""), "")
                or source
            )
            resume = (
                row.get(headers.get("resume pdf", ""), "")
                or row.get(headers.get("resume", ""), "")
                or row.get(headers.get("resume to use", ""), "")
                or row.get(headers.get("resume to attach", ""), "")
            )
            resume = resolve_resume_path(resume, path)
            category = "Opportunity report"
            if "recent_interns" in path.as_posix():
                category = "Recent internship pass"
            elif "selected_internships" in path.as_posix():
                category = "Selected internship batch"
            elif "next_applications" in path.as_posix() or path.name == "next_application_materials.md":
                category = "Next application batch"
            elif "email_leads" in path.as_posix() or path.name == "email_lead_application_materials.md":
                category = "Email lead batch"
            elif "social_leads" in path.as_posix() or path.name == "social_lead_application_materials.md":
                category = "Social lead batch"
            status = ""
            if heading.lower() == "roles skipped":
                category = "Skipped / low-fit role"
                status = "Skip"
            upsert_opportunity(
                opps,
                parsed_company,
                parsed_role,
                path,
                Status=status,
                Priority=priority,
                Category=category,
                Fit=fit,
                Location__Constraint=location,
                **{
                    "Location / Constraint": location,
                    "Apply Route": apply_route,
                    "Apply Link / Email": extract_links_emails(apply_route or source) or apply_route,
                    "Resume PDF": resume,
                    "Notes": source if source and source != apply_route else "",
                },
            )


def company_from_heading(heading: str) -> str:
    heading = clean_md(heading)
    heading = re.sub(r"^\d+\.\s*", "", heading)
    heading = re.sub(
        r"\s+(Email|Note|DM|Form|Comment|Cover Note|Keka Cover Note|Wellfound Note|LinkedIn Comment|WhatsApp)$",
        "",
        heading,
        flags=re.I,
    )
    if heading.lower().startswith("proofbuild"):
        return "proof_of_Build"
    return canonical_company(heading)


def parse_draft_files(opps: dict[str, dict]) -> set[str]:
    companies_with_drafts = set()
    ignore_headings = {
        "source notes",
        "short follow-up after 3-4 days",
        "apply links and priority",
        "apply links and sources",
        "priority",
        "certifiedaijobs note",
    }
    for path in DRAFT_FILES:
        if not path.exists():
            continue
        current = ""
        current_lines: list[str] = []

        def flush():
            if not current or current.lower() in ignore_headings:
                return
            current_lower = current.lower()
            counts_as_applied = not (
                "form" in current_lower
                or current_lower in {"walmart note", "navi form note", "navi note"}
            )
            company = company_from_heading(current)
            if not company:
                return
            companies_with_drafts.add(norm_key(company))
            chunk = "\n".join(current_lines)
            route_lines = [
                line
                for line in current_lines[:12]
                if re.search(
                    r"\b(To|Send to|Recommended route|Best route|Apply|Founder route|Safer hiring route|Email fallback|Routing fallback|Cold-routing fallback)\b",
                    line,
                    flags=re.I,
                )
            ]
            emails_links = extract_links_emails("\n".join(route_lines) or "\n".join(current_lines[:8]))
            matching = [
                row
                for row in opps.values()
                if norm_key(row["Company"]) == norm_key(company)
                or norm_key(company) in norm_key(row["Company"])
                or norm_key(row["Company"]) in norm_key(company)
            ]
            if not matching:
                matching = [
                    upsert_opportunity(
                        opps,
                        company,
                        "General / careers target",
                        path,
                        Status="Applied by draft" if counts_as_applied else "Open",
                        Category="Cold outreach target",
                    )
                ]
            for row in matching:
                if counts_as_applied:
                    row["Status"] = "Applied by draft"
                row["Materials / Draft File"] = append_unique(row["Materials / Draft File"], rel(path))
                if emails_links:
                    row["Apply Link / Email"] = append_unique(row["Apply Link / Email"], emails_links)
                row["Source Files"] = append_unique(row["Source Files"], rel(path))

        for line in path.read_text(encoding="utf-8", errors="replace").splitlines():
            match = re.match(r"^##\s+(.+)$", line)
            if match:
                flush()
                current = clean_md(match.group(1))
                current_lines = []
            elif current:
                current_lines.append(line)
        flush()
    return companies_with_drafts


def attach_active_pdfs(opps: dict[str, dict]) -> None:
    active_dir = ROOT / "active_email_attachments_2026_05_26"
    for pdf in sorted(active_dir.glob("*.pdf")):
        stem_norm = norm_key(pdf.stem)
        best = None
        best_score = 0
        for row in opps.values():
            company_score = len(norm_key(row["Company"])) if norm_key(row["Company"]) in stem_norm else 0
            role_score = len(norm_key(row["Role"])) if norm_key(row["Role"]) in stem_norm else 0
            score = company_score + role_score
            for token in re.findall(r"[A-Za-z0-9]+", row["Company"]):
                token_norm = norm_key(token)
                if len(token_norm) > 3 and token_norm in stem_norm:
                    score += len(token_norm) + 10
            if row["Company"] == "CredResolve / Banking Stack" and "credresolve" in stem_norm:
                score += 100
            if score > best_score:
                best, best_score = row, score
        if best and best_score:
            best["Attachment PDF"] = append_unique(best["Attachment PDF"], rel(pdf))
            best["Status"] = "Applied by draft"
            best["Source Files"] = append_unique(best["Source Files"], rel(pdf))


def finalize_opportunities(opps: dict[str, dict]) -> list[dict]:
    for row in opps.values():
        if row["Status"] == "Unclassified":
            row["Status"] = "Open"
        lower_notes = (row["Notes"] + " " + row["Fit"]).lower()
        if "no longer accepting" in lower_notes or "posting unavailable" in lower_notes:
            if row["Status"] != "Applied by draft":
                row["Status"] = "Closed / inactive"
        if "do not pay" in lower_notes or "paid membership" in lower_notes:
            row["Status"] = "Skip"
        if not row["Apply Route"] and row["Apply Link / Email"]:
            row["Apply Route"] = row["Apply Link / Email"]
        if not row["Next Action"]:
            if row["Status"] == "Open":
                if "form" in (row["Company"] + " " + row["Apply Route"] + " " + row["Notes"]).lower():
                    row["Next Action"] = "Submit or confirm form submission."
                else:
                    row["Next Action"] = "Verify posting/current route, then apply."
            elif row["Status"] == "Applied by draft":
                row["Next Action"] = "Track response; follow up after 3-5 days if relevant."
            elif row["Status"] in {"Needs verification", "Closed / inactive"}:
                row["Next Action"] = "Verify status before spending more time."
            elif row["Status"] in {"Eliminated", "Skip"}:
                row["Next Action"] = "Do not pursue unless circumstances change."
        if not row["Last Seen"]:
            row["Last Seen"] = "2026-05-27"
    status_order = {
        "Open": 0,
        "Needs verification": 1,
        "Applied by draft": 2,
        "Closed / inactive": 3,
        "Eliminated": 4,
        "Skip": 5,
    }
    priority_order = {"P0": 0, "P0 if eligible": 1, "P1": 2, "P1 if eligible": 3, "P2": 4, "P2 if eligible": 5}
    return sorted(
        opps.values(),
        key=lambda r: (
            status_order.get(r["Status"], 9),
            priority_order.get(r["Priority"], 20),
            norm_key(r["Company"]),
            norm_key(r["Role"]),
        ),
    )


def parse_contacts() -> list[dict]:
    rows = []
    seen = set()
    for path in CONTACT_FILES:
        if not path.exists():
            continue
        current = ""
        for idx, line in enumerate(path.read_text(encoding="utf-8", errors="replace").splitlines(), start=1):
            heading = re.match(r"^#{1,3}\s+(.+)$", line)
            if heading:
                current = clean_md(heading.group(1))
            emails = EMAIL_RE.findall(line)
            if not emails:
                continue
            company = company_from_heading(current) if current else ""
            for email in emails:
                if email == "ishuvam.pal@gmail.com":
                    continue
                key = (company, email, rel(path))
                if key in seen:
                    continue
                seen.add(key)
                rows.append(
                    {
                        "Company": company,
                        "Contact / Route": current,
                        "Email / URL": email,
                        "Contact Type": "Email",
                        "Source File": f"{rel(path)}:{idx}",
                        "Notes": clean_md(line),
                    }
                )
    rows.sort(key=lambda r: (norm_key(r["Company"]), r["Email / URL"], r["Source File"]))
    return rows


def parse_company_targets() -> list[dict]:
    rows: list[dict] = []

    # Structured recheck table.
    recheck = ROOT / "outreach_category_recheck_2026-05-11.md"
    if recheck.exists():
        for _, row in iter_markdown_tables(recheck):
            company = row.get("Company", "")
            if not company:
                continue
            rows.append(
                {
                    "Company": canonical_company(company),
                    "Track": "Outreach category recheck",
                    "Recommendation / Tier": row.get("Current category", ""),
                    "Rank": "",
                    "Fit Area": "",
                    "Source File": rel(recheck),
                    "Notes": f"{row.get('Current read', '')} {row.get('Application note and sources', '')}".strip(),
                }
            )

    for path in [ROOT / "startup_outreach_tiers.md", ROOT / "startup_backend_infra_ai_tiers.md"]:
        if not path.exists():
            continue
        tier = ""
        for line in path.read_text(encoding="utf-8", errors="replace").splitlines():
            h2 = re.match(r"^##\s+(.+)$", line)
            if h2:
                tier = clean_md(h2.group(1))
                continue
            bullet = re.match(r"^-\s+([^:]+):\s+(.+)$", line)
            if not bullet or "Source" in tier or "Best Shortlist" in tier:
                continue
            rows.append(
                {
                    "Company": canonical_company(bullet.group(1)),
                    "Track": path.stem.replace("_", " ").title(),
                    "Recommendation / Tier": tier,
                    "Rank": "",
                    "Fit Area": "",
                    "Source File": rel(path),
                    "Notes": clean_md(bullet.group(2)),
                }
            )

    rankings = ROOT / "backend_infra_ai_rankings.md"
    if rankings.exists():
        track = ""
        rank = ""
        pending_company = ""
        for line in rankings.read_text(encoding="utf-8", errors="replace").splitlines():
            h2 = re.match(r"^##\s+(.+?)\s+Ranking$", line)
            if h2:
                track = h2.group(1)
                continue
            numbered = re.match(r"^(\d+)\.\s+(.+)$", line)
            if numbered and track:
                rank, pending_company = numbered.group(1), canonical_company(numbered.group(2))
                continue
            reason = re.match(r"^Reason:\s+(.+)$", line)
            if reason and pending_company:
                rows.append(
                    {
                        "Company": pending_company,
                        "Track": f"{track} ranking",
                        "Recommendation / Tier": "Top 10",
                        "Rank": rank,
                        "Fit Area": track,
                        "Source File": rel(rankings),
                        "Notes": clean_md(reason.group(1)),
                    }
                )
                pending_company = ""
            bullet = re.match(r"^-\s+(.+)$", line)
            if bullet and track:
                rows.append(
                    {
                        "Company": canonical_company(bullet.group(1)),
                        "Track": f"{track} ranking",
                        "Recommendation / Tier": "Follow-up / shortlist",
                        "Rank": "",
                        "Fit Area": track,
                        "Source File": rel(rankings),
                        "Notes": "",
                    }
                )

    rows.sort(key=lambda r: (norm_key(r["Company"]), r["Track"], str(r["Rank"])))
    return rows


def parse_resume_inventory(opps: list[dict]) -> list[dict]:
    pdfs = sorted([*ROOT.glob("tailored_resumes/**/*.pdf"), *ROOT.glob("active_email_attachments_2026_05_26/*.pdf")])
    rows = []
    for pdf in pdfs:
        stem = pdf.stem
        stem_norm = norm_key(stem)
        best = None
        best_score = 0
        for opp in opps:
            c_norm = norm_key(opp["Company"])
            r_norm = norm_key(opp["Role"])
            score = 0
            if c_norm and c_norm in stem_norm:
                score += len(c_norm) + 50
            for token in re.findall(r"[A-Za-z0-9]+", opp["Company"]):
                if len(token) > 3 and norm_key(token) in stem_norm:
                    score += len(token)
            if r_norm and r_norm in stem_norm:
                score += min(len(r_norm), 50)
            if score > best_score:
                best, best_score = opp, score
        mapped = f"{best['Company']} - {best['Role']}" if best and best_score else ""
        rows.append(
            {
                "File": rel(pdf),
                "Collection": rel(pdf.parent),
                "Inferred Company": best["Company"] if best and best_score else "",
                "Mapped Opportunity": mapped,
                "Size KB": round(pdf.stat().st_size / 1024, 1),
                "Modified": datetime.fromtimestamp(pdf.stat().st_mtime).strftime("%Y-%m-%d %H:%M"),
            }
        )
    return rows


def source_index() -> list[dict]:
    important = set()
    for path in [STATUS_FILE, *OPPORTUNITY_TABLE_FILES, *DRAFT_FILES, *CONTACT_FILES, *TARGET_RESEARCH_FILES]:
        if path.exists():
            important.add(path)
    rows = []
    for path in sorted(important):
        if not path.exists():
            continue
        parsed_as = []
        if path == STATUS_FILE:
            parsed_as.append("authoritative status")
        if path in OPPORTUNITY_TABLE_FILES:
            parsed_as.append("opportunity/enrichment table")
        if path in DRAFT_FILES:
            parsed_as.append("draft/application evidence")
        if path in CONTACT_FILES:
            parsed_as.append("contact extraction")
        if path in TARGET_RESEARCH_FILES:
            parsed_as.append("company research")
        rows.append(
            {
                "Source File": rel(path),
                "Parsed As": ", ".join(parsed_as),
                "Modified": datetime.fromtimestamp(path.stat().st_mtime).strftime("%Y-%m-%d %H:%M"),
                "Size KB": round(path.stat().st_size / 1024, 1),
                "Notes": "",
            }
        )
    return rows


def write_csv(rows: list[dict], path: Path) -> None:
    if not rows:
        return
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0].keys()))
        writer.writeheader()
        writer.writerows(rows)


def rows_to_sheet(ws, rows: list[dict], table_name: str, freeze: str = "A2") -> None:
    if not rows:
        ws.append(["No data"])
        return
    headers = list(rows[0].keys())
    ws.append(headers)
    for row in rows:
        ws.append([row.get(header, "") for header in headers])
    style_sheet(ws, table_name, headers)
    ws.freeze_panes = freeze


def style_sheet(ws, table_name: str, headers: list[str]) -> None:
    ws.sheet_view.showGridLines = False
    max_row = ws.max_row
    max_col = ws.max_column
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER
    for row_idx in range(2, max_row + 1):
        fill = EVEN_FILL if row_idx % 2 == 0 else ODD_FILL
        for col_idx in range(1, max_col + 1):
            cell = ws.cell(row_idx, col_idx)
            cell.fill = fill
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = THIN_BORDER
            if isinstance(cell.value, str):
                if cell.value.startswith("http"):
                    cell.hyperlink = cell.value.split(" | ")[0]
                    cell.style = "Hyperlink"
                elif re.fullmatch(EMAIL_RE, cell.value.split(" | ")[0] or ""):
                    cell.hyperlink = f"mailto:{cell.value.split(' | ')[0]}"
                    cell.style = "Hyperlink"
    widths = {}
    for col_idx, header in enumerate(headers, start=1):
        values = [str(ws.cell(row_idx, col_idx).value or "") for row_idx in range(1, min(max_row, 60) + 1)]
        max_len = max([len(header), *[min(len(v), 80) for v in values]])
        if header in {"Notes", "Fit", "Apply Route", "Apply Link / Email", "Source Files", "Materials / Draft File"}:
            width = min(max(max_len, 22), 55)
        elif header in {"Company", "Role", "Mapped Opportunity"}:
            width = min(max(max_len, 18), 38)
        else:
            width = min(max(max_len, 12), 28)
        widths[col_idx] = width
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.auto_filter.ref = ws.dimensions
    if max_row >= 2 and max_col >= 1:
        safe_name = re.sub(r"[^A-Za-z0-9_]", "", table_name)[:25]
        table = Table(displayName=safe_name, ref=ws.dimensions)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=False,
            showColumnStripes=False,
        )
        try:
            ws.add_table(table)
        except ValueError:
            pass

    if "Status" in headers:
        status_col = get_column_letter(headers.index("Status") + 1)
        for status, color in STATUS_FILLS.items():
            ws.conditional_formatting.add(
                f"A2:{get_column_letter(max_col)}{max_row}",
                FormulaRule(
                    formula=[f'${status_col}2="{status}"'],
                    fill=PatternFill("solid", fgColor=color),
                ),
            )
    if "Priority" in headers:
        priority_col = get_column_letter(headers.index("Priority") + 1)
        for priority, color in PRIORITY_FILLS.items():
            ws.conditional_formatting.add(
                f"{priority_col}2:{priority_col}{max_row}",
                FormulaRule(
                    formula=[f'${priority_col}2="{priority}"'],
                    fill=PatternFill("solid", fgColor=color),
                ),
            )


def write_dashboard(wb: Workbook, opportunities: list[dict], targets: list[dict], contacts: list[dict], resumes: list[dict]) -> None:
    ws = wb.create_sheet("Dashboard", 0)
    ws.sheet_view.showGridLines = False
    ws["A1"] = "Internship Tracker"
    ws["A1"].font = Font(size=20, bold=True, color="111827")
    ws["A2"] = "Generated from local intern directory on 2026-05-27. Drafted cold mail/application material counts as applied/drafted."
    ws["A2"].alignment = Alignment(wrap_text=True)

    status_counts = Counter(row["Status"] for row in opportunities)
    category_counts = Counter(row["Category"] or "Uncategorized" for row in opportunities)
    priority_open = Counter(row["Priority"] or "No priority" for row in opportunities if row["Status"] == "Open")

    blocks = [
        ("Status", status_counts),
        ("Open By Priority", priority_open),
        ("Category", category_counts),
    ]
    start_col = 1
    for title, counter in blocks:
        ws.cell(4, start_col, title).fill = HEADER_FILL
        ws.cell(4, start_col, title).font = Font(color="FFFFFF", bold=True)
        ws.cell(4, start_col + 1, "Count").fill = HEADER_FILL
        ws.cell(4, start_col + 1).font = Font(color="FFFFFF", bold=True)
        for idx, (name, count) in enumerate(counter.most_common(), start=5):
            ws.cell(idx, start_col, name)
            ws.cell(idx, start_col + 1, count)
            for col in [start_col, start_col + 1]:
                ws.cell(idx, col).fill = EVEN_FILL if idx % 2 == 0 else ODD_FILL
                ws.cell(idx, col).border = THIN_BORDER
        start_col += 3

    summary = [
        ("Total opportunity rows", len(opportunities)),
        ("Open rows", status_counts.get("Open", 0)),
        ("Applied/drafted rows", status_counts.get("Applied by draft", 0)),
        ("Needs verification / inactive / skip rows", sum(v for k, v in status_counts.items() if k != "Open" and k != "Applied by draft")),
        ("Company research rows", len(targets)),
        ("Contact rows", len(contacts)),
        ("Resume PDFs indexed", len(resumes)),
    ]
    ws["A15"] = "Quick Counts"
    ws["A15"].fill = HEADER_FILL
    ws["A15"].font = Font(color="FFFFFF", bold=True)
    ws["B15"].fill = HEADER_FILL
    for idx, (label, value) in enumerate(summary, start=16):
        ws.cell(idx, 1, label)
        ws.cell(idx, 2, value)
        ws.cell(idx, 1).fill = EVEN_FILL if idx % 2 == 0 else ODD_FILL
        ws.cell(idx, 2).fill = EVEN_FILL if idx % 2 == 0 else ODD_FILL
        ws.cell(idx, 1).border = THIN_BORDER
        ws.cell(idx, 2).border = THIN_BORDER

    ws["D15"] = "Highest-signal open queue"
    ws["D15"].fill = HEADER_FILL
    ws["D15"].font = Font(color="FFFFFF", bold=True)
    ws["E15"].fill = HEADER_FILL
    open_rows = [r for r in opportunities if r["Status"] == "Open"][:15]
    for idx, row in enumerate(open_rows, start=16):
        ws.cell(idx, 4, f"{row['Company']} - {row['Role']}")
        ws.cell(idx, 5, row["Priority"] or row["Next Action"])
        for col in [4, 5]:
            ws.cell(idx, col).fill = EVEN_FILL if idx % 2 == 0 else ODD_FILL
            ws.cell(idx, col).alignment = Alignment(wrap_text=True, vertical="top")
            ws.cell(idx, col).border = THIN_BORDER

    for col, width in {"A": 34, "B": 12, "D": 58, "E": 28, "G": 32, "H": 12}.items():
        ws.column_dimensions[col].width = width
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A4"


def main() -> None:
    opps = parse_status_file()
    parse_opportunity_tables(opps)
    parse_draft_files(opps)
    attach_active_pdfs(opps)
    opportunities = finalize_opportunities(opps)
    contacts = parse_contacts()
    targets = parse_company_targets()
    resumes = parse_resume_inventory(opportunities)
    sources = source_index()

    write_csv(opportunities, OUT_CSV)

    wb = Workbook()
    default = wb.active
    wb.remove(default)

    write_dashboard(wb, opportunities, targets, contacts, resumes)
    rows_to_sheet(wb.create_sheet("Opportunities"), opportunities, "Opportunities")
    rows_to_sheet(wb.create_sheet("Open Queue"), [r for r in opportunities if r["Status"] == "Open"], "OpenQueue")
    rows_to_sheet(wb.create_sheet("Applied or Drafted"), [r for r in opportunities if r["Status"] == "Applied by draft"], "AppliedDrafted")
    rows_to_sheet(
        wb.create_sheet("Needs Verification"),
        [r for r in opportunities if r["Status"] not in {"Open", "Applied by draft"}],
        "NeedsVerification",
    )
    rows_to_sheet(wb.create_sheet("Contacts"), contacts, "Contacts")
    rows_to_sheet(wb.create_sheet("Company Targets"), targets, "CompanyTargets")
    rows_to_sheet(wb.create_sheet("Resume Inventory"), resumes, "ResumeInventory")
    rows_to_sheet(wb.create_sheet("Source Index"), sources, "SourceIndex")

    wb.save(OUT_XLSX)
    print(f"Wrote {OUT_XLSX}")
    print(f"Wrote {OUT_CSV}")
    print(f"Opportunities: {len(opportunities)}")
    print(f"Open: {sum(1 for r in opportunities if r['Status'] == 'Open')}")
    print(f"Applied by draft: {sum(1 for r in opportunities if r['Status'] == 'Applied by draft')}")
    print(f"Contacts: {len(contacts)}")
    print(f"Company targets: {len(targets)}")
    print(f"Resume PDFs: {len(resumes)}")


if __name__ == "__main__":
    main()
